import pandas as pd
import openpyxl
import re
import streamlit as st

class ExcelReader:
    def __init__(self, file):
        self.file = file
        self.workbook = openpyxl.load_workbook(file, data_only=True)
        self.sheet_names = self.workbook.sheetnames

    def get_sheets(self):
        return self.sheet_names

    def read_sheet(self, sheet_name):
        sheet = self.workbook[sheet_name]
        data = []
        for row in sheet.iter_rows(values_only=True):
            data.append(row)
        return pd.DataFrame(data)

class FeedbackParser:
    def __init__(self):
        # Regex to match: Name: 1) Answer1 2) Answer2 3) Answer3
        # Supports various formats and multi-line answers
        self.pattern = re.compile(
            r"^(?P<name>.*?):\s*"
            r"1\)\s*(?P<q1>.*?)\s*"
            r"2\)\s*(?P<q2>.*?)\s*"
            r"3\)\s*(?P<q3>.*)$",
            re.DOTALL | re.MULTILINE
        )

    def parse_cell(self, text):
        if not isinstance(text, str) or ":" not in text:
            return None
        
        match = self.pattern.match(text.strip())
        if match:
            return {
                "name": match.group("name").strip(),
                "q1": match.group("q1").strip(),
                "q2": match.group("q2").strip(),
                "q3": match.group("q3").strip()
            }
        
        # Fallback if the standard pattern doesn't match perfectly
        # (Handle case where 1) might be missing but : is present)
        parts = text.split(":", 1)
        name = parts[0].strip()
        rest = parts[1].strip()
        
        # Try splitting by the question markers
        q1_match = re.search(r"1\)\s*(.*?)\s*(?=2\)|$)", rest, re.DOTALL)
        q2_match = re.search(r"2\)\s*(.*?)\s*(?=3\)|$)", rest, re.DOTALL)
        q3_match = re.search(r"3\)\s*(.*)$", rest, re.DOTALL)
        
        return {
            "name": name,
            "q1": q1_match.group(1).strip() if q1_match else "",
            "q2": q2_match.group(1).strip() if q2_match else "",
            "q3": q3_match.group(1).strip() if q3_match else ""
        }

def process_feedback_file(file):
    reader = ExcelReader(file)
    parser = FeedbackParser()
    all_records = []

    for sheet_name in reader.get_sheets():
        df = reader.read_sheet(sheet_name)
        # Module name is typically the sheet name
        module_name = sheet_name
        
        # We assume the first column might contain Video names and others contain responses
        # Or each table in the sheet represents a video. 
        # Based on user description: "separated by modules and the videos in each module"
        # We'll treat columns as videos if there are multiple, or rows.
        # Let's iterate through all cells and if it matches the pattern, we add it.
        # We'll need to identify which video it belongs to.
        # Simplified assumption: If a cell is found, we look at the header for Video name.
        
        headers = df.iloc[0]
        for row_idx in range(1, len(df)):
            for col_idx in range(len(df.columns)):
                cell_value = df.iloc[row_idx, col_idx]
                parsed = parser.parse_cell(cell_value)
                if parsed:
                    video_name = headers[col_idx] if col_idx < len(headers) else f"Video {col_idx}"
                    parsed["module"] = module_name
                    parsed["video"] = video_name
                    all_records.append(parsed)
    
    return all_records

if __name__ == "__main__":
    st.set_page_config(page_title="Feedback Compiler", layout="wide")
    st.title("📑 Qualitative Feedback Compiler")
    st.write("Extract and organize qualitative feedback from your Excel module trackers.")

    uploaded_file = st.file_uploader("Choose an Excel file", type="xlsx")

    if uploaded_file:
        try:
            records = process_feedback_file(uploaded_file)
            if records:
                st.success(f"Successfully parsed {len(records)} feedback entries!")
                st.session_state['records'] = records
            else:
                st.warning("No feedback entries found matching the expected format (Name: 1) ... 2) ... 3) ...)")
        except Exception as e:
            st.error(f"Error processing file: {e}")

    if 'records' in st.session_state:
        records = st.session_state['records']
        # Placeholders for the views
        view_mode = st.radio("Select View Mode", ["Compiled View", "Thematic View"])
        
        if view_mode == "Compiled View":
            st.header("Module & Video Feedback")
            
            # Grouping by Module then Video
            modules = sorted(list(set(r['module'] for r in records)))
            selected_module = st.selectbox("Select Module", modules)
            
            module_records = [r for r in records if r['module'] == selected_module]
            videos = sorted(list(set(r['video'] for r in module_records)))
            
            for video in videos:
                with st.expander(f"🎬 {video}", expanded=True):
                    video_records = [r for r in module_records if r['video'] == video]
                    for r in video_records:
                        st.subheader(f"👤 {r['name']}")
                        col1, col2, col3 = st.columns(3)
                        with col1:
                            st.info("**Q1: Incorrect/Needs Update?**")
                            st.write(r['q1'])
                        with col2:
                            st.info("**Q2: Missing/Want More?**")
                            st.write(r['q2'])
                        with col3:
                            st.info("**Q3: Other Feedback?**")
                            st.write(r['q3'])
                        st.divider()
        else:
            st.header("Thematic Question Analysis")
            question_choice = st.selectbox(
                "Select Question to Analyze",
                [
                    "Q1: Is anything incorrect or needs updating?",
                    "Q2: Is anything missing, or is there something you’d like more of?",
                    "Q3: Any other feedback or suggestions?"
                ]
            )
            
            q_key = "q1" if "Q1" in question_choice else "q2" if "Q2" in question_choice else "q3"
            
            for r in records:
                if r[q_key].strip():
                    with st.container():
                        st.markdown(f"**{r['name']}** ({r['module']} > {r['video']})")
                        st.write(r[q_key])
                        st.divider()
